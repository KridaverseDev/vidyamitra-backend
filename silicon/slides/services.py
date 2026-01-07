from django.db import IntegrityError, transaction
from django.shortcuts import get_object_or_404
from ninja.errors import HttpError
from pptx.util import Pt

from silicon.knowledge.models import Course, Document
from silicon.knowledge.services import DocumentService
from silicon.slides import schemas
from silicon.slides.ai.llm import SlideGeneratorLLM
from silicon.slides.presentation import GeneratePresentation

from .models import Favorite, Presentation, Slides

# Define custom formatting options
TITLE_FONT_SIZE = Pt(30)
SLIDE_FONT_SIZE = Pt(16)


class PresentationService:
    def __init__(self, user=None):
        """
        Initialize PresentationService with optional user for API key resolution.
        
        Args:
            user: CustomUser instance (optional) - for user-specific API keys
        """
        self.user = user
        self.llm = SlideGeneratorLLM(user=user)  # Pass user for API key resolution
        self.document = DocumentService(user=user)  # Pass user for API key resolution

    def get_presentation(self, id: int):
        """Get the presentation, along with its corresponding slides."""

        slides = []
        ppt = get_object_or_404(Presentation, pk=id)
        for slide in Slides.objects.filter(presentation_id=id):
            slides.append(slide)
        return {"presentation": ppt, "slides": slides}

    def generate_presentation(
        self, request, document_id: int, payload: schemas.GeneratePPTIn
    ):
        doc = get_object_or_404(Document, pk=document_id)
        course = get_object_or_404(Course, pk=doc.course_id)
        user = request.auth

        # fetching the query for the given description from the document.
        prompt = f"""{payload.query}"""
        context = self.document.multi_query(namespace=doc.namespace, query=prompt)
        response = self.llm.generate_presentation(
            payload.title, payload.count + 2, context["response"]
        )

        print(f"{response=}")

        with transaction.atomic():
            # Creating a presentation object.
            ppt = Presentation.objects.create(
                title=payload.title, query=payload.query, course=course, user=user
            )

            for index, slide in enumerate(response["slides"], start=1):
                title = slide["title"]
                content = slide["content"]

                items = "\n".join([f"- {item}" for item in content])

                # Create a slide object.
                Slides.objects.create(title=title, content=items, presentation=ppt)

        raise HttpError(201, "Presentation is created successfully")

    def delete_presentation(self, id: int):
        """Delete a presentation object."""
        ppt = get_object_or_404(Presentation, pk=id)
        ppt.delete()

    def update_slide(self, slide_id: int, payload: schemas.SlideIn):
        """Update the contents of the given slide."""
        updated_fields = payload.dict(exclude_unset=True)
        slide = get_object_or_404(Slides, pk=slide_id)
        try:
            for attr, value in updated_fields.items():
                setattr(slide, attr, value)
            slide.save()
        except Exception as e:
            raise HttpError(500, str(e))

    def export_presentation(self, id: int):
        """Export the presentation into a presentation .pptx file and return S3 URL."""
        titles = []
        contents = []
        ppt = get_object_or_404(Presentation, pk=id)
        try:
            for slide in Slides.objects.filter(presentation_id=id):
                titles.append(slide.title)
                contents.append(slide.content)
            
            if not titles:
                raise HttpError(400, "No slides found in presentation")
            
            result = self._create_presentation(
                ppt.title, slide_titles=titles, slide_contents=contents
            )
            return result
        except HttpError:
            raise
        except Exception as e:
            raise HttpError(500, f"Error exporting presentation: {str(e)}")

    def _create_presentation(self, prs_name, slide_titles, slide_contents):
        """Create a presentation file using pptx and return S3 file info."""
        print(f"{len(slide_titles)=}")
        print(f"{len(slide_contents)=}")

        p = GeneratePresentation()
        p.new(prs_name)

        for idx, (title, content) in enumerate(zip(slide_titles, slide_contents)):
            print(f"---- slide[{idx + 1}]: {title} ----")
            p.new_slide(title, content)

        # Save to S3 and return file info
        result = p.save(prs_name)
        return result

    def favorite_presentation(self, request, id: int):
        """Mark a presentation as favorite."""
        ppt = get_object_or_404(Presentation, pk=id)
        try:
            Favorite.objects.create(ppt=ppt, user=request.auth)
        except IntegrityError:
            raise HttpError(409, "Already marked as favorite")
        raise HttpError(201, "Presentation is marked as favorite")
